"""CSV / JSON 解析、去重、聚合与 Excel 导出测试。"""

import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fscan_parser import export_workbook  # noqa: E402
from fscan_parser.merge import build_assets, dedupe, merge_web_records  # noqa: E402
from fscan_parser.models import RecordType  # noqa: E402
from fscan_parser.parsers import parse_file, parse_text, sniff_format  # noqa: E402

FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")


def test_parse_csv():
    records, stats = parse_file(os.path.join(FIXTURES, "result.csv"))
    assert stats.formats.get("csv") == 12
    assert len(records) == 12
    services = [r for r in records if r.type is RecordType.SERVICE]
    assert services[0].service == "ssh"
    assert services[2].title == "Apache Tomcat/9.0"


def test_parse_json():
    records, stats = parse_file(os.path.join(FIXTURES, "result.json"))
    assert stats.formats.get("json") == 12
    vulns = [r for r in records if r.type is RecordType.VULN]
    assert any(v.credential == "root/123456" for v in vulns)
    assert all(r.time for r in records)


def test_parse_ndjson_realtime():
    path = os.path.join(FIXTURES, "result.json.realtime.tmp")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write('{"type":"PORT","target":"10.0.0.9","status":"open","details":{"port":8080}}\n')
        fh.write('{"type":"VULN","target":"http://10.0.0.9:8080","status":"vulnerable",'
                 '"details":{"vulnerability_name":"CVE-2021-44228","type":"vulnerability"}}\n')
    try:
        records, stats = parse_file(path)
        assert stats.formats.get("json") == 2
        assert records[0].port == 8080
        assert records[1].vuln_name == "CVE-2021-44228"
    finally:
        os.remove(path)


def test_sniff_format():
    assert sniff_format("a.txt", "# ===== 存活主机 =====") == "txt"
    assert sniff_format("a.csv", "# Hosts\nTarget\n") == "csv"
    assert sniff_format("x", '{"hosts": []}') == "json"
    assert sniff_format("y", "[+] 1.1.1.1:80 open") == "txt"


def test_dedupe():
    r1, _ = parse_text("# ===== 开放端口 =====\n1.1.1.1:80\n1.1.1.1:80\n")
    assert len(r1) == 2
    kept, dropped = dedupe(r1)
    assert len(kept) == 1
    assert dropped == 1


def test_merge_web_records():
    with open(os.path.join(FIXTURES, "result.txt"), encoding="utf-8") as fh:
        records, _ = parse_text(fh.read())
    merged = merge_web_records(records)
    assert all(r.type is not RecordType.WEB for r in merged)
    assert len(merged) == len(records)


def test_build_assets():
    with open(os.path.join(FIXTURES, "result.txt"), encoding="utf-8") as fh:
        records, _ = parse_text(fh.read())
    assets = {a.host: a for a in build_assets(records)}
    assert set(assets) == {"192.168.1.1", "192.168.1.5"}
    assert assets["192.168.1.1"].open_ports == {22, 445}
    assert assets["192.168.1.5"].open_ports == {80, 443, 3389}
    assert assets["192.168.1.5"].highest_severity == "高危"
    assert "http://192.168.1.5:80" in assets["192.168.1.5"].web_urls


def test_export_workbook(tmp_path):
    with open(os.path.join(FIXTURES, "result.txt"), encoding="utf-8") as fh:
        records, stats = parse_text(fh.read(), path="result.txt")
    out = tmp_path / "out.xlsx"
    export_workbook(records, str(out), stats=stats, source_files=["result.txt"])
    assert out.exists() and out.stat().st_size > 0

    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames == [
        "概览",
        "资产视图",
        "存活主机",
        "开放端口",
        "服务信息",
        "Web服务",
        "WebTitle",
        "漏洞信息",
        "未知记录",
    ]
    assert wb["存活主机"].max_row == 3
    assert wb["开放端口"].max_row == 4
    assert wb["服务信息"].max_row == 5
    assert wb["漏洞信息"].max_row == 4
    # 表头冻结与自动筛选
    assert wb["开放端口"].freeze_panes == "A2"
    assert wb["开放端口"].auto_filter.ref.startswith("A1:")
    # 端口列应是数字而非文本
    assert isinstance(wb["开放端口"].cell(row=2, column=2).value, int)


def test_export_prevents_formula_injection(tmp_path):
    payload = '# ===== 漏洞信息 =====\n=http://evil/x "=cmd|calc"\n'
    records, _ = parse_text(payload)
    out = tmp_path / "inj.xlsx"
    export_workbook(records, str(out))
    wb = openpyxl.load_workbook(str(out))
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    assert not value.startswith("=")
