"""Excel 导出：把解析结果写成结构化的 .xlsx 工作簿。

工作表结构：
    概览     扫描统计、Top 服务、Top 漏洞、目标分布
    资产视图 一台主机一行（端口列表 / 服务 / Web / 漏洞数 / 最高等级）
    存活主机
    开放端口
    服务信息
    Web服务
    WebTitle 旧版 [*] WebTitle / [*] 网站标题 行（1.8.x、2.0.0 才有）
    漏洞信息
    未知记录 解析器无法归类但保留了原文的行（便于人工复核）

实现上使用 openpyxl 的 write_only 模式：大文件（数万条记录）下比逐格访问快约 3~4 倍，
同时通过模块级样式单例复用样式对象，避免每个单元格都新建一份样式。
"""

from __future__ import annotations

import json
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .merge import HostAsset, build_assets, sort_records
from .models import ParseStats, Record, RecordType, Severity
from .targets import guess_scheme

# ---------------------------------------------------------------- 安全处理
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_EXCEL_CELL_LIMIT = 32767


def _safe(value: Any) -> Any:
    """防止公式注入：以 = + - @ 开头的文本前置单引号，并截断超长内容。"""
    if not isinstance(value, str):
        return value
    if len(value) > _EXCEL_CELL_LIMIT:
        value = value[:_EXCEL_CELL_LIMIT] + "..."
    if value.startswith(_FORMULA_PREFIXES) and len(value) > 1:
        return "'" + value
    return value


# ---------------------------------------------------------------- 样式单例
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
BODY_BOLD = Font(name="微软雅黑", size=10, bold=True)
TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="1F3864")
SECTION_FONT = Font(name="微软雅黑", size=12, bold=True, color="1F3864")
SUB_FONT = Font(name="微软雅黑", size=11, bold=True, color="1F3864")
SUBTLE_FONT = Font(name="微软雅黑", size=9, color="808080")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ZEBRA_FILL = PatternFill("solid", fgColor="F2F5FA")
CENTER = Alignment(horizontal="center", vertical="center")
TOP = Alignment(vertical="top")
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
LEFT = Alignment(horizontal="left")

# 预组合的「样式包」：openpyxl 每次给单元格赋样式都要做一次哈希去重，
# 复用同几个组合可把这一步从数十秒降到几秒。
_PLAIN = {"font": BODY_FONT, "border": BORDER, "alignment": TOP}
_PLAIN_ZEBRA = {"font": BODY_FONT, "border": BORDER, "alignment": TOP, "fill": ZEBRA_FILL}
_WRAP = {"font": BODY_FONT, "border": BORDER, "alignment": TOP_WRAP}
_WRAP_ZEBRA = {"font": BODY_FONT, "border": BORDER, "alignment": TOP_WRAP, "fill": ZEBRA_FILL}
_HEADER = {"font": HEADER_FONT, "border": BORDER, "alignment": CENTER, "fill": HEADER_FILL}



SEVERITY_FILLS = {
    Severity.HIGH: PatternFill("solid", fgColor="FFC7CE"),
    Severity.MEDIUM: PatternFill("solid", fgColor="FFEB9C"),
    Severity.LOW: PatternFill("solid", fgColor="DDEBF7"),
    Severity.INFO: PatternFill("solid", fgColor="EDEDED"),
}
SEVERITY_FONTS = {
    Severity.HIGH: Font(name="微软雅黑", size=10, bold=True, color="9C0006"),
    Severity.MEDIUM: Font(name="微软雅黑", size=10, color="9C5700"),
    Severity.LOW: Font(name="微软雅黑", size=10, color="1F4E78"),
    Severity.INFO: Font(name="微软雅黑", size=10, color="595959"),
}

_STYLE_CACHE: Dict[tuple, Any] = {}


def _style_obj(style: Dict[str, Any], ws: Worksheet) -> Any:
    """把样式包固化成 openpyxl 的 _style 对象，供所有单元格复用。

    直接给 font/fill/border 赋值会触发 openpyxl 的不可变样式去重（每次都要哈希），
    大文件下这是主要开销；复用同一个 _style 对象可把写入时间减半。
    """
    workbook_key = id(ws.parent)
    key = (workbook_key, id(style))
    obj = _STYLE_CACHE.get(key)
    if obj is None:
        probe = WriteOnlyCell(ws, value="")
        probe.font = style["font"]
        probe.border = style["border"]
        probe.alignment = style["alignment"]
        if "fill" in style:
            probe.fill = style["fill"]
        obj = probe._style
        _STYLE_CACHE[key] = obj
    return obj


def _styled(ws: Worksheet, value: Any, style: Dict[str, Any]) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=_safe(value))
    cell._style = _style_obj(style, ws)
    return cell


_SEVERITY_STYLES = {
    sev: {
        "font": SEVERITY_FONTS[sev],
        "border": BORDER,
        "alignment": CENTER,
        "fill": SEVERITY_FILLS[sev],
    }
    for sev in Severity
}

_KNOWN_TYPES = {
    RecordType.HOST,
    RecordType.PORT,
    RecordType.SERVICE,
    RecordType.VULN,
    RecordType.WEBTITLE,
}

# 列定义: (表头, 取值函数, 列宽)
Column = tuple[str, Any, int]

SHEET_COLUMNS: Dict[str, List[Column]] = {
    "存活主机": [
        ("IP/主机", lambda r: r.target.host, 22),
        ("状态", lambda r: r.status or "alive", 10),
        ("来源文件", lambda r: _short_path(r.source), 28),
        ("行号", lambda r: r.line_no, 8),
    ],
    "开放端口": [
        ("IP/主机", lambda r: r.target.host, 22),
        ("端口", lambda r: r.port, 8),
        ("状态", lambda r: r.status or "open", 10),
        ("来源文件", lambda r: _short_path(r.source), 28),
        ("行号", lambda r: r.line_no, 8),
    ],
    "服务信息": [
        ("IP/主机", lambda r: r.target.host, 22),
        ("端口", lambda r: r.port, 8),
        ("服务", lambda r: r.service, 14),
        ("版本", lambda r: r.version, 26),
        ("HTTP状态", lambda r: r.details.get("status", ""), 10),
        ("Server头", lambda r: r.details.get("server", ""), 24),
        ("标题", lambda r: r.title, 36),
        ("指纹", lambda r: ",".join(r.fingerprint_list), 26),
        ("Banner", lambda r: r.banner, 46),
        ("产品", lambda r: r.details.get("product", ""), 18),
        ("操作系统", lambda r: r.details.get("os", ""), 20),
        ("来源文件", lambda r: _short_path(r.source), 28),
        ("行号", lambda r: r.line_no, 8),
    ],
    "Web服务": [
        ("URL", lambda r: _web_url(r), 46),
        ("IP/主机", lambda r: r.target.host, 22),
        ("端口", lambda r: r.port, 8),
        ("协议", lambda r: r.target.scheme or guess_scheme(r.target.raw, r.service, r.port), 8),
        ("HTTP状态", lambda r: r.details.get("status", ""), 10),
        ("标题", lambda r: r.title, 36),
        ("Server头", lambda r: r.details.get("server", ""), 24),
        ("指纹", lambda r: ",".join(r.fingerprint_list), 26),
        ("来源文件", lambda r: _short_path(r.source), 28),
        ("行号", lambda r: r.line_no, 8),
    ],
    "WebTitle": [
        ("URL", lambda r: r.details.get("url") or _web_url(r), 46),
        ("IP/主机", lambda r: r.target.host, 22),
        ("端口", lambda r: r.port, 8),
        ("协议", lambda r: r.target.scheme or guess_scheme(r.target.raw, r.service, r.port), 8),
        ("HTTP状态", lambda r: r.details.get("status", ""), 10),
        ("长度", lambda r: r.details.get("length", ""), 8),
        ("标题", lambda r: r.title, 44),
        ("跳转", lambda r: r.details.get("redirect", ""), 34),
        ("来源文件", lambda r: _short_path(r.source), 28),
        ("行号", lambda r: r.line_no, 8),
    ],
    "漏洞信息": [
        ("等级", lambda r: r.severity.value, 8),
        ("IP/主机", lambda r: r.target.host, 22),
        ("端口", lambda r: r.port, 8),
        ("目标", lambda r: r.target.raw, 40),
        ("漏洞/类型", lambda r: r.vuln_name, 40),
        ("服务", lambda r: r.service, 14),
        ("账号", lambda r: r.details.get("username", ""), 16),
        ("密码", lambda r: r.details.get("password", ""), 16),
        ("详情/描述", lambda r: _vuln_desc(r), 52),
        ("来源文件", lambda r: _short_path(r.source), 28),
        ("行号", lambda r: r.line_no, 8),
    ],
    "未知记录": [
        ("类型", lambda r: r.type.value, 10),
        ("目标", lambda r: r.target.raw, 36),
        ("状态", lambda r: r.status, 16),
        ("原始内容", lambda r: r.raw, 70),
        ("来源文件", lambda r: _short_path(r.source), 28),
        ("行号", lambda r: r.line_no, 8),
    ],
}


def export_workbook(
    records: List[Record],
    output_path: str,
    stats: Optional[ParseStats] = None,
    source_files: Optional[List[str]] = None,
    title: str = "fscan 扫描结果",
) -> str:
    """写出 xlsx 文件，返回输出路径。"""
    _STYLE_CACHE.clear()  # 样式索引与工作簿绑定，每次导出重新建立
    records = sort_records(records)
    assets = build_assets(records)

    wb = Workbook(write_only=True)
    _write_summary(wb, records, assets, stats, source_files, title)
    _write_assets(wb, assets)
    _write_detail_sheet(wb, "存活主机", [r for r in records if r.type is RecordType.HOST], SHEET_COLUMNS["存活主机"])
    _write_detail_sheet(wb, "开放端口", [r for r in records if r.type is RecordType.PORT], SHEET_COLUMNS["开放端口"])
    _write_detail_sheet(wb, "服务信息", [r for r in records if r.type is RecordType.SERVICE], SHEET_COLUMNS["服务信息"])
    _write_detail_sheet(wb, "Web服务", [r for r in records if _is_web(r)], SHEET_COLUMNS["Web服务"])
    _write_detail_sheet(
        wb, "WebTitle", [r for r in records if r.type is RecordType.WEBTITLE], SHEET_COLUMNS["WebTitle"]
    )
    _write_detail_sheet(wb, "漏洞信息", [r for r in records if r.type is RecordType.VULN], SHEET_COLUMNS["漏洞信息"])
    _write_detail_sheet(
        wb, "未知记录", [r for r in records if r.type not in _KNOWN_TYPES], SHEET_COLUMNS["未知记录"]
    )

    wb.save(output_path)
    return output_path


# ------------------------------------------------------------------ helper
def _cell(ws: Worksheet, value: Any, font: Font = BODY_FONT, fill: Optional[PatternFill] = None,
          align: Optional[Alignment] = None, border: Optional[Border] = None) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=_safe(value))
    cell.font = font
    if fill is not None:
        cell.fill = fill
    if align is not None:
        cell.alignment = align
    if border is not None:
        cell.border = border
    return cell


def _blank(ws: Worksheet) -> WriteOnlyCell:
    return WriteOnlyCell(ws, value=None)


def _pad(row: List[WriteOnlyCell], width: int, ws: Worksheet) -> List[WriteOnlyCell]:
    while len(row) < width:
        row.append(_blank(ws))
    return row


# ------------------------------------------------------------------ 概览
def _write_summary(
    wb: Workbook,
    records: List[Record],
    assets: List[HostAsset],
    stats: Optional[ParseStats],
    source_files: Optional[List[str]],
    title: str,
) -> None:
    ws = wb.create_sheet("概览")
    ws.sheet_view.showGridLines = False
    for col, width in (("A", 3), ("B", 26), ("C", 18), ("D", 6), ("E", 24), ("F", 10)):
        ws.column_dimensions[col].width = width

    vulns = [r for r in records if r.type is RecordType.VULN]
    rows: List[List[WriteOnlyCell]] = [
        [],  # 第 1 行留白
        [_blank(ws), _cell(ws, title, font=TITLE_FONT)],
        [_blank(ws), _cell(ws, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", font=SUBTLE_FONT)],
        [],
        [_blank(ws), _cell(ws, "扫描统计", font=SECTION_FONT)],
    ]

    overview = [
        ("存活主机", len([r for r in records if r.type is RecordType.HOST])),
        ("开放端口", len([r for r in records if r.type is RecordType.PORT])),
        ("服务记录", len([r for r in records if r.type is RecordType.SERVICE])),
        ("Web 服务", len([r for r in records if _is_web(r)])),
        ("WebTitle", len([r for r in records if r.type is RecordType.WEBTITLE])),
        ("漏洞记录", len(vulns)),
        ("高危漏洞", len([v for v in vulns if v.severity is Severity.HIGH])),
        ("资产主机数", len(assets)),
    ]
    for label, value in overview:
        rows.append(
            [_blank(ws), _cell(ws, label), _cell(ws, value, font=BODY_BOLD, align=LEFT)]
        )

    # fscan 输出格式不带版本号，这里显示自动识别（或 --version-spec 指定）的结果
    detection = getattr(stats, "detection", None) if stats is not None else None
    if detection is not None or (stats is not None and stats.versions):
        rows.append([])
        rows.append([_blank(ws), _cell(ws, "fscan 版本", font=SECTION_FONT)])
        era_counts = dict(stats.versions) if stats is not None else {}
        label = _era_label(era_counts) or (detection.label if detection is not None else "-")
        rows.append([_blank(ws), _cell(ws, "版本区间"), _cell(ws, label)])
        if detection is not None:
            source = "手动指定" if getattr(detection, "forced", False) else "自动识别"
            # 混合版本时第一个文件不一定代表多数，标注清楚归属
            scope = "" if len(era_counts) <= 1 else f"（样本来自 {detection.era.value}）"
            rows.append(
                [_blank(ws), _cell(ws, "置信度"), _cell(ws, f"{round(detection.confidence * 100)}%（{source}）{scope}")]
            )
            for i, line in enumerate(getattr(detection, "evidence_lines", lambda: [])()):
                rows.append(
                    [_blank(ws), _cell(ws, "命中特征" if i == 0 else ""), _cell(ws, line, font=SUBTLE_FONT)]
                )
            keywords = ("同时命中", "无法识别", "未匹配", "推测")
            for note in getattr(detection, "notes", []):
                if any(k in note for k in keywords):
                    rows.append([_blank(ws), _cell(ws, "提示"), _cell(ws, note, font=SUBTLE_FONT)])

    if stats is not None:
        rows.append([])
        rows.append([_blank(ws), _cell(ws, "解析信息", font=SECTION_FONT)])
        for label, value in (
            ("输入文件数", stats.files),
            ("忽略行数", stats.skipped),
            ("解析异常", stats.failed),
            ("识别格式", ", ".join(f"{k}×{v}" for k, v in stats.formats.items()) or "-"),
        ):
            rows.append([_blank(ws), _cell(ws, label), _cell(ws, value)])

    if source_files:
        rows.append([])
        rows.append([_blank(ws), _cell(ws, "来源文件", font=SECTION_FONT)])
        for path in source_files:
            rows.append([_blank(ws), _cell(ws, path, font=SUBTLE_FONT)])

    rows.append([])

    # 四张统计小表并排：服务 / 漏洞 / 端口 / 网段
    services = _counter(r.service for r in records if r.type is RecordType.SERVICE and r.service)
    vuln_top = _counter(v.vuln_name for v in vulns if v.vuln_name)
    ports = _counter(str(r.port) for r in records if r.type in (RecordType.PORT, RecordType.SERVICE) and r.port)
    segments = _counter(_segment(a.host) for a in assets)

    tables = [
        (1, "服务分布", services, 12),   # B 列
        (4, "漏洞 TOP10", vuln_top, 10),  # E 列
    ]
    tables2 = [
        (1, "端口分布 TOP15", ports, 15),
        (4, "网段分布", segments, 15),
    ]
    for block in (tables, tables2):
        rows.extend(_counter_rows(ws, block))
        rows.append([])

    for row in rows:
        ws.append(_pad(row, 6, ws))


def _counter_rows(ws: Worksheet, block) -> List[List[WriteOnlyCell]]:
    """把并排的多张小表按行对齐成 WriteOnly 行序列。"""
    max_len = max(len(items[:limit]) for _offset, _title, items, limit in block)
    out: List[List[WriteOnlyCell]] = []

    title_row: List[WriteOnlyCell] = []
    header_row: List[WriteOnlyCell] = []
    for offset, title, _items, _limit in block:
        while len(title_row) < offset:
            title_row.append(_blank(ws))
        title_row.append(_cell(ws, title, font=SUB_FONT))
        while len(header_row) < offset:
            header_row.append(_blank(ws))
        header_row.append(_cell(ws, "名称", font=HEADER_FONT, fill=HEADER_FILL, align=CENTER, border=BORDER))
        header_row.append(_cell(ws, "数量", font=HEADER_FONT, fill=HEADER_FILL, align=CENTER, border=BORDER))
    out.append(title_row)
    out.append(header_row)

    for i in range(max_len):
        line: List[WriteOnlyCell] = []
        for offset, _title, items, limit in block:
            while len(line) < offset:
                line.append(_blank(ws))
            pair = items[i] if i < min(len(items), limit) else None
            line.append(_cell(ws, pair[0] if pair else "", border=BORDER))
            line.append(_cell(ws, pair[1] if pair else "", border=BORDER))
        out.append(line)
    return out


def _era_label(counts: Dict[str, int]) -> str:
    """把「版本区间 → 文件数」拼成一行展示文本；单版本时直接给区间名。"""
    from .version import FscanEra

    if not counts:
        return ""
    label_of = {e.value: e.label for e in FscanEra}
    if len(counts) == 1:
        era = next(iter(counts))
        return label_of.get(era, era)
    parts = [f"{label_of.get(k, k)}×{v} 个文件" for k, v in sorted(counts.items(), key=lambda kv: -kv[1])]
    return "混合版本：" + "，".join(parts)


def _counter(values: Iterable[str]) -> List[tuple[str, int]]:
    counts: Dict[str, int] = {}
    for v in values:
        if not v:
            continue
        counts[v] = counts.get(v, 0) + 1
    return sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))


def _segment(host: str) -> str:
    parts = host.split(".")
    if len(parts) == 4 and all(p.isdigit() for p in parts):
        return ".".join(parts[:3]) + ".0/24"
    return "其他"


# ------------------------------------------------------------------ 资产视图
ASSET_COLUMNS: List[Column] = [
    ("IP/主机", lambda a: a.host, 22),
    ("存活", lambda a: "是" if a.is_alive else "否", 8),
    ("开放端口", lambda a: a.port_list, 40),
    ("端口数", lambda a: len(a.open_ports), 8),
    ("服务", lambda a: ", ".join(a.service_set), 30),
    ("Web地址", lambda a: "\n".join(a.web_urls), 50),
    ("Web数", lambda a: len(a.web_urls), 8),
    ("漏洞数", lambda a: len(a.vulns), 8),
    ("最高等级", lambda a: a.highest_severity, 10),
    ("漏洞列表", lambda a: "\n".join(sorted({v.vuln_name for v in a.vulns if v.vuln_name})), 50),
]


def _write_assets(wb: Workbook, assets: List[HostAsset]) -> None:
    ws = wb.create_sheet("资产视图")
    _write_header(ws, ASSET_COLUMNS)

    for i, asset in enumerate(assets, start=2):
        zebra = i % 2 == 0
        row: List[WriteOnlyCell] = []
        for j, (_header, getter, _width) in enumerate(ASSET_COLUMNS, start=1):
            wrap = j in (3, 5, 6, 10)
            style = (_WRAP_ZEBRA if zebra else _WRAP) if wrap else (_PLAIN_ZEBRA if zebra else _PLAIN)
            row.append(_styled(ws, getter(asset), style))
        sev_enum = next((s for s in Severity if s.value == asset.highest_severity), None)
        if sev_enum is not None:
            row[8] = _styled(ws, asset.highest_severity, _SEVERITY_STYLES[sev_enum])
        ws.append(row)
    _finalize_sheet(ws, len(assets), len(ASSET_COLUMNS))


# ------------------------------------------------------------------ 明细表
def _write_detail_sheet(wb: Workbook, name: str, records: List[Record], columns: List[Column]) -> None:
    ws = wb.create_sheet(name)
    _write_header(ws, columns)

    severity_col = name == "漏洞信息"
    for i, rec in enumerate(records, start=2):
        zebra = i % 2 == 0
        row: List[WriteOnlyCell] = []
        for j, (_header, getter, width) in enumerate(columns, start=1):
            try:
                value = getter(rec)
            except Exception:
                value = ""
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            if j == 1 and severity_col:
                row.append(_styled(ws, value, _SEVERITY_STYLES[rec.severity]))
            elif width >= 30:
                row.append(_styled(ws, value, _WRAP_ZEBRA if zebra else _WRAP))
            else:
                row.append(_styled(ws, value, _PLAIN_ZEBRA if zebra else _PLAIN))
        ws.append(row)
    _finalize_sheet(ws, len(records), len(columns))


def _write_header(ws: Worksheet, columns: List[Column]) -> None:
    # write_only 模式下视图属性必须在首次 append 前设置
    ws.freeze_panes = "A2"
    ws.sheet_view.zoomScale = 100
    for idx, (header, _getter, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 22
    ws.append([_styled(ws, header, _HEADER) for header, _g, _w in columns])


def _finalize_sheet(ws: Worksheet, data_rows: int, n_cols: int) -> None:
    if data_rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{data_rows + 1}"


# ------------------------------------------------------------------ 工具
def _short_path(path: str) -> str:
    if not path:
        return ""
    return path.replace("\\", "/").rsplit("/", 1)[-1]


def _is_web(rec: Record) -> bool:
    if rec.type is not RecordType.SERVICE:
        return False
    return bool(
        rec.details.get("is_web")
        or rec.details.get("title")
        or rec.details.get("status")
        or rec.details.get("server")
        or rec.service.lower() in ("http", "https")
    )


def _web_url(rec: Record) -> str:
    if rec.target.url:
        return rec.target.url
    scheme = guess_scheme(rec.target.raw, rec.service, rec.port)
    return f"{scheme}://{rec.target.endpoint}"


def _vuln_desc(rec: Record) -> str:
    for key in ("description", "parameters", "info"):
        val = rec.details.get(key)
        if val:
            if isinstance(val, (dict, list)):
                return json.dumps(val, ensure_ascii=False)
            text = str(val)
            # 多行原文里首行是结论，后面的才是细节
            lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
            if len(lines) > 1:
                return " | ".join(lines[1:])
            return text
    refs = rec.details.get("references")
    if isinstance(refs, list) and refs:
        return " | ".join(str(r) for r in refs)
    return ""
